"""Photo alt text.

Reported from a real issue: none of its 8 pictures had a description,
and the email gave 6 of its 7 photos alt text like `Picture 122515128`
-- what a recipient sees when images are blocked, and what a screen
reader reads aloud. `_drawing_to_img` used `pic:cNvPr@descr` or, failing
that, `pic:cNvPr@name`: the object name Word generates, which describes
nothing. And Word's Alt Text box writes to `wp:docPr@descr`, which was
never read, so a description an editor did type could be ignored.

The rule pinned here:
  1. a real description wins -- `wp:docPr@descr`, else `pic:cNvPr@descr`;
  2. an object name is never alt text;
  3. a picture Word marks decorative keeps `alt=""`;
  4. otherwise the alt text is the title of the section the picture is
     in, numbered "(1)", "(2)"... when that section has several such
     pictures;
  5. the dean's photo keeps its own alt text.
"""

from __future__ import annotations

import re
from html import unescape
from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

from scripts.config import MERIDIAN_TEMPLATE
from scripts.docx_parser import (
    BodyParagraph,
    BulletList,
    TableBlock,
    paragraph_to_html,
    parse,
)
from scripts.renderer import _resolve_media, attach_image_urls, render
from scripts.webapp import build_from_bytes

_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_ADEC_NS = "http://schemas.microsoft.com/office/drawing/2017/decorative"
# The `a:ext` uri Word writes around the decorative flag.
_DECORATIVE_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"

DEAN_ALT = "Dean of the Graduate School of Medicine"


# ---------- helpers ----------
@pytest.fixture
def png(tmp_path: Path) -> Path:
    path = tmp_path / "photo.png"
    Image.new("RGB", (40, 20), "white").save(path)
    return path


def _picture(run, png: Path, *, descr: str | None = None,
             cnv_descr: str | None = None, name: str | None = None,
             cnv_name: str | None = None,
             decorative: str | None = None) -> None:
    """Add a picture to `run`, then set the attributes Word would write.

    python-docx itself names the picture `Picture N` on `wp:docPr` and
    puts the source file's name (`photo.png`) on `pic:cNvPr` -- so an
    untouched call already exercises "object names are not alt text".
    """
    inline = run.add_picture(str(png))._inline
    doc_pr = inline.find(qn("wp:docPr"))
    cnv_pr = inline.find(".//" + qn("pic:cNvPr"))
    for element, attr, value in ((doc_pr, "descr", descr),
                                 (cnv_pr, "descr", cnv_descr),
                                 (doc_pr, "name", name),
                                 (cnv_pr, "name", cnv_name)):
        if value is not None:
            element.set(attr, value)
    if decorative is not None:
        ext_lst = etree.SubElement(doc_pr, f"{{{_A_NS}}}extLst")
        ext = etree.SubElement(ext_lst, f"{{{_A_NS}}}ext", uri=_DECORATIVE_URI)
        etree.SubElement(ext, f"{{{_ADEC_NS}}}decorative", val=decorative)


def _photo_paragraph(d, png: Path, **attrs) -> None:
    _picture(d.add_paragraph().add_run(), png, **attrs)


def _strict_doc():
    """A template-shaped document: its first table is the masthead, which
    the strict parser always skips."""
    d = docx.Document()
    d.add_table(rows=1, cols=2)
    return d


def _save(d, tmp_path: Path, name: str = "issue.docx") -> Path:
    path = tmp_path / name
    d.save(str(path))
    return path


_IMG = re.compile(r"<img\b[^>]*>")
_ALT = re.compile(r'\balt="([^"]*)"')


def _img_alts(html: str) -> list[str]:
    """The alt text of every picture in `html`, unescaped, in order.
    A picture with no alt attribute at all shows up as None."""
    out = []
    for tag in _IMG.findall(html):
        m = _ALT.search(tag)
        out.append(unescape(m.group(1)) if m else None)
    return out


def _block_htmls(block) -> list[str]:
    if isinstance(block, BodyParagraph):
        return [block.html]
    if isinstance(block, BulletList):
        return list(block.items)
    if isinstance(block, TableBlock):
        return [cell for row in block.rows for cell in row]
    return []


def _alts_by_section(path: Path) -> list[list[str]]:
    return [
        [alt for block in section.blocks for html in _block_htmls(block)
         for alt in _img_alts(html)]
        for section in parse(path).sections
    ]


# ---------- 1. a real description wins ----------
def test_the_description_from_words_alt_text_box_is_used(tmp_path, png):
    """Word's Alt Text box stores the description on `wp:docPr@descr`."""
    d = _strict_doc()
    d.add_paragraph("1. Research")
    _photo_paragraph(d, png, descr="The lab team at the award ceremony")
    assert _alts_by_section(_save(d, tmp_path)) == [
        ["The lab team at the award ceremony"]]


def test_a_description_on_the_picture_properties_is_used(tmp_path, png):
    d = _strict_doc()
    d.add_paragraph("1. Research")
    _photo_paragraph(d, png, cnv_descr="Poster session")
    assert _alts_by_section(_save(d, tmp_path)) == [["Poster session"]]


def test_the_alt_text_box_wins_over_the_picture_properties(tmp_path, png):
    d = _strict_doc()
    d.add_paragraph("1. Research")
    _photo_paragraph(d, png, descr="From the Alt Text box",
                     cnv_descr="Older description")
    assert _alts_by_section(_save(d, tmp_path)) == [["From the Alt Text box"]]


def test_a_description_is_stripped_and_a_blank_one_does_not_count(
        tmp_path, png):
    d = _strict_doc()
    d.add_paragraph("1. Research")
    _photo_paragraph(d, png, descr="  Poster session \n")
    _photo_paragraph(d, png, descr=" \n ", cnv_descr="\t")
    assert _alts_by_section(_save(d, tmp_path)) == [
        ["Poster session", "Research"]]


# ---------- 2. object names are never descriptions ----------
@pytest.mark.parametrize("name", [
    "Picture 122515128", "drawing", "図 3", "image1.png", "photo.png"])
def test_an_object_name_is_never_alt_text(tmp_path, png, name):
    d = _strict_doc()
    d.add_paragraph("1. Research")
    _photo_paragraph(d, png, name=name, cnv_name=name)
    path = _save(d, tmp_path)
    assert _alts_by_section(path) == [["Research"]]


# ---------- 3. decorative pictures stay silent ----------
@pytest.mark.parametrize("val", ["1", "true"])
def test_a_decorative_picture_keeps_an_empty_alt(tmp_path, png, val):
    """Word's "Mark as decorative" means a screen reader should skip the
    picture. Filling in the section title would undo the editor's choice
    -- even where a hand-edited file still carries a description."""
    d = _strict_doc()
    d.add_paragraph("1. Research")
    _photo_paragraph(d, png, decorative=val)
    _photo_paragraph(d, png, decorative=val, descr="Left-over description")
    assert _alts_by_section(_save(d, tmp_path)) == [["", ""]]


def test_decorative_switched_off_is_not_decorative(tmp_path, png):
    d = _strict_doc()
    d.add_paragraph("1. Research")
    _photo_paragraph(d, png, decorative="0")
    assert _alts_by_section(_save(d, tmp_path)) == [["Research"]]


# ---------- 4. the section title ----------
def test_a_single_undescribed_photo_gets_the_section_title(tmp_path, png):
    d = _strict_doc()
    d.add_paragraph("1. Message from the Dean")
    d.add_paragraph("Welcome to the first issue.")
    _photo_paragraph(d, png)
    assert _alts_by_section(_save(d, tmp_path)) == [["Message from the Dean"]]


def test_several_undescribed_photos_are_numbered_in_document_order(
        tmp_path, png):
    """Numbering restarts in each section, and a sub-heading does not
    restart it: the strict parser's unit is the numbered section."""
    d = _strict_doc()
    d.add_paragraph("1. Research & Academic Updates")
    both = d.add_paragraph()
    _picture(both.add_run(), png)
    _picture(both.add_run(), png)
    d.add_heading("Awards", level=2)
    _photo_paragraph(d, png)
    d.add_paragraph("2. International Collaboration")
    _photo_paragraph(d, png)
    assert _alts_by_section(_save(d, tmp_path)) == [
        ["Research & Academic Updates (1)", "Research & Academic Updates (2)",
         "Research & Academic Updates (3)"],
        ["International Collaboration"],
    ]


def test_only_undescribed_photos_are_counted(tmp_path, png):
    d = _strict_doc()
    d.add_paragraph("1. News")
    _photo_paragraph(d, png, descr="A caption")
    _photo_paragraph(d, png)
    _photo_paragraph(d, png, decorative="1")
    _photo_paragraph(d, png)
    d.add_paragraph("2. Events")
    _photo_paragraph(d, png, descr="Opening ceremony")
    _photo_paragraph(d, png)
    assert _alts_by_section(_save(d, tmp_path)) == [
        ["A caption", "News (1)", "", "News (2)"],
        ["Opening ceremony", "Events"],
    ]


def test_photos_in_bullets_and_table_cells_are_titled(tmp_path, png):
    """The Featured Highlights cards are tables, so a table cell is where
    many photos live."""
    d = _strict_doc()
    d.add_paragraph("1. Featured Highlights")
    _photo_paragraph(d, png)
    _picture(d.add_paragraph("", style="List Paragraph").add_run(), png)
    cards = d.add_table(rows=1, cols=3)
    _picture(cards.cell(0, 0).paragraphs[0].add_run(), png)
    _picture(cards.cell(0, 2).paragraphs[0].add_run(), png)
    assert _alts_by_section(_save(d, tmp_path)) == [
        [f"Featured Highlights ({n})" for n in (1, 2, 3, 4)]]


def test_the_title_is_escaped_for_the_attribute(tmp_path, png):
    d = _strict_doc()
    d.add_paragraph('1. Q&A <live> "now"')
    _photo_paragraph(d, png)
    _photo_paragraph(d, png)
    path = _save(d, tmp_path)
    html = " ".join(h for s in parse(path).sections for b in s.blocks
                    for h in _block_htmls(b))
    assert 'alt="Q&amp;A &lt;live&gt; &quot;now&quot; (1)"' in html
    assert _alts_by_section(path) == [
        ['Q&A <live> "now" (1)', 'Q&A <live> "now" (2)']]
    # The tag is still whole: the escaped title did not end the attribute.
    for tag in _IMG.findall(html):
        assert tag.endswith("/>") and "style=" in tag


def test_lenient_parse_uses_the_nearest_preceding_heading(tmp_path, png):
    """A Word file without numbered sections is one synthetic section
    called "Newsletter". Its Word headings are what a reader sees as the
    sections, so they name the photos; numbering restarts under each."""
    d = docx.Document()
    d.add_paragraph("Welcome to this newsletter.")
    _photo_paragraph(d, png)
    d.add_heading("Recent News", level=1)
    _photo_paragraph(d, png)
    _photo_paragraph(d, png, descr="Our new building")
    _photo_paragraph(d, png)
    d.add_heading("Awards", level=2)
    box = d.add_table(rows=1, cols=1)
    _picture(box.cell(0, 0).paragraphs[0].add_run(), png)
    path = _save(d, tmp_path)
    assert [s.title for s in parse(path).sections] == ["Newsletter"]
    assert _alts_by_section(path) == [[
        "Newsletter",
        "Recent News (1)", "Our new building", "Recent News (2)",
        "Awards",
    ]]


def test_lenient_parse_numbers_photos_before_any_heading(tmp_path, png):
    d = docx.Document()
    d.add_paragraph("No headings at all.")
    _photo_paragraph(d, png)
    _photo_paragraph(d, png)
    assert _alts_by_section(_save(d, tmp_path)) == [
        ["Newsletter (1)", "Newsletter (2)"]]


def test_a_paragraph_converted_on_its_own_still_has_an_alt(png):
    """`paragraph_to_html` is public and knows nothing of sections. What
    it returns must still be a valid picture -- an `<img>` with no alt
    at all has its file name read out by a screen reader."""
    d = docx.Document()
    p = d.add_paragraph()
    _picture(p.add_run(), png)
    assert _img_alts(paragraph_to_html(p)) == [""]


def test_the_renderer_never_lets_the_untitled_marker_through(png):
    """`parse` fills the marker in; HTML built any other way still must
    not carry it to a recipient."""
    d = docx.Document()
    p = d.add_paragraph()
    _picture(p.add_run(), png)
    out = _resolve_media(paragraph_to_html(p), {})
    assert "data-meridian" not in out
    assert _img_alts(out) == [""]


# ---------- 5. the dean's photo ----------
@pytest.mark.parametrize("name", [
    "Nagoya_university_school_medicine_dean.jpg",
    "NAGOYA_UNIVERSITY_SCHOOL_MEDICINE_DEAN.JPG",
    "Nagoya_university_school_medicine_dean",
])
def test_the_dean_photo_keeps_its_alt_text(tmp_path, png, name):
    """The template builder inserts the photo with python-docx, which
    records the source file name as the object name -- the only trace of
    which photo it is, since the media part is renamed `imageN.jpg`. The
    photo is not counted among the section's undescribed photos."""
    d = _strict_doc()
    d.add_paragraph("1. Message from the Dean")
    _photo_paragraph(d, png, cnv_name=name)
    _photo_paragraph(d, png)
    html = render(attach_image_urls(parse(_save(d, tmp_path)), {}))
    media = [t for t in _IMG.findall(html) if "media://" in t]
    assert [_img_alts(t)[0] for t in media] == [
        DEAN_ALT, "Message from the Dean"]


def test_a_description_on_the_dean_photo_still_wins(tmp_path, png):
    d = _strict_doc()
    d.add_paragraph("1. Message from the Dean")
    _photo_paragraph(d, png, cnv_name="Nagoya_university_school_medicine_dean.jpg",
                     descr="Prof. Katsuno in his office")
    assert _alts_by_section(_save(d, tmp_path)) == [
        ["Prof. Katsuno in his office"]]


@pytest.mark.skipif(not MERIDIAN_TEMPLATE.exists(),
                    reason="MERIDIAN template not built yet")
def test_the_templates_dean_photo_alt_is_unchanged_end_to_end():
    result = build_from_bytes(MERIDIAN_TEMPLATE.read_bytes(), issue=0)
    for doc in (result.html, result.standalone_html):
        alts = _img_alts(doc)
        assert alts.count(DEAN_ALT) == 1, alts
        assert not any("dean" in (alt or "").lower() and alt != DEAN_ALT
                       for alt in alts), alts


# ---------- end to end ----------
def test_alt_text_reaches_the_mail_and_the_standalone_html(tmp_path, png):
    """Through the entry point the browser build calls: parse, render,
    inline, CID resolution, and the self-contained download."""
    d = _strict_doc()
    d.add_paragraph("1. Research & Academic Updates")
    _photo_paragraph(d, png, name="Picture 122515128",
                     cnv_name="Picture 122515128")
    _photo_paragraph(d, png)
    d.add_paragraph("2. Contact Information")
    _photo_paragraph(d, png, descr="Campus map")
    result = build_from_bytes(_save(d, tmp_path).read_bytes(), issue=7)
    for doc in (result.html, result.standalone_html):
        alts = _img_alts(doc)
        assert "Research & Academic Updates (1)" in alts, alts
        assert "Research & Academic Updates (2)" in alts, alts
        assert "Campus map" in alts, alts
        assert None not in alts, "a picture lost its alt attribute"
        assert not any((alt or "").startswith("Picture ") for alt in alts)
        assert "data-meridian" not in doc, "an internal marker leaked"

"""The website builds the same newsletter as the desktop toolkit.

The page runs the real `scripts/` package inside Pyodide, with its own
copies of the libraries loaded from `web/pyodide/<version>/`. Nothing in
the rest of the suite executes that runtime -- so a broken or missing
wheel, a dependency the page forgot to name, a runtime the CDN stopped
serving, or a library version that renders differently would reach
editors before any test noticed.

These run the actual engine under Node: the vendored runtime, the
packages `web/app.js` loads and the committed bundle. They compare its
email with the desktop build of the same Word files. They need Node and
the vendored runtime, so they run only with `MERIDIAN_WEB_ENGINE=1`, which
the `web-engine` workflow sets:

    python web/vendor_pyodide.py
    MERIDIAN_WEB_ENGINE=1 python -m pytest tests/test_web_engine.py
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
from importlib.metadata import version
from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

from scripts.webapp import build_from_bytes

REPO_ROOT = Path(__file__).resolve().parent.parent
RUNNER = REPO_ROOT / "tests" / "web_engine" / "run.cjs"

pytestmark = pytest.mark.skipif(
    os.environ.get("MERIDIAN_WEB_ENGINE") != "1",
    reason="set MERIDIAN_WEB_ENGINE=1 to run the website engine "
           "(needs Node and `python web/vendor_pyodide.py`)")


def _fixture_docx(path: Path) -> Path:
    """A small newsletter touching what the engine has to get right:
    numbered sections, justified and centred text, a photo, a bullet list
    and a table."""
    png = path.with_suffix(".png")
    Image.new("RGB", (80, 40), (0, 63, 136)).save(png)
    d = docx.Document()
    d.add_table(rows=1, cols=2)                    # masthead
    d.add_paragraph("1. News")
    justified = d.add_paragraph(
        "A justified paragraph, long enough to wrap in the email.")
    justified.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    photo = d.add_paragraph()
    photo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo.add_run().add_picture(str(png))
    d.add_paragraph("A bullet item.", style="List Paragraph")
    table = d.add_table(rows=2, cols=2)
    for r in range(2):
        for c in range(2):
            table.cell(r, c).text = f"Cell {r}{c}."
    d.add_paragraph("2. Events")
    d.add_paragraph("See you at the seminar on Friday.")
    d.save(str(path))
    return path


@pytest.fixture(scope="module")
def engine(tmp_path_factory):
    node = shutil.which("node")
    assert node, "MERIDIAN_WEB_ENGINE=1, but Node is not installed"
    version = json.loads((REPO_ROOT / "web" / "pyodide-assets.json").read_text(
        encoding="utf-8"))["pyodide_version"]
    assert (REPO_ROOT / "web" / "pyodide" / version / "pyodide.js").exists(), (
        f"run `python web/vendor_pyodide.py` first (no Pyodide {version} "
        f"in web/pyodide/)")
    work = tmp_path_factory.mktemp("engine")
    documents = {
        "template": REPO_ROOT / "Meridian_Newsletter_Template.docx",
        "fixture": _fixture_docx(work / "fixture.docx"),
    }
    output = work / "engine.json"
    run = subprocess.run(
        [node, str(RUNNER), str(REPO_ROOT / "web"), str(output),
         *map(str, documents.values())],
        capture_output=True, text=True, encoding="utf-8", errors="replace",
        timeout=600)
    assert run.returncode == 0, run.stdout + run.stderr
    built = json.loads(output.read_text(encoding="utf-8"))
    return {name: (path, built[path.name]) for name, path in documents.items()}


@pytest.mark.parametrize("name", ["template", "fixture"])
def test_the_website_builds_the_same_email_as_the_desktop(engine, name):
    path, browser = engine[name]
    desktop = build_from_bytes(path.read_bytes(), issue=7)
    assert browser["subject"] == desktop.subject
    assert browser["errors"] == list(desktop.errors)
    assert browser["section_count"] == desktop.section_count
    assert browser["photo_count"] == desktop.photo_count
    assert browser["plaintext"] == desktop.plaintext
    assert browser["html"] == desktop.html, (
        "the website's email HTML differs from the desktop build's; the "
        f"website runs {json.dumps(browser['versions'])}")


def test_the_website_runs_the_desktop_versions_of_the_pinned_libraries(engine):
    """Only exact pins are compared. A `>=` requirement lets the desktop
    install a newer release than the runtime ships, which is expected; if
    that ever changes the email, the HTML comparison above says so."""
    _, browser = engine["fixture"]
    for project in ("css-inline", "beautifulsoup4", "python-docx"):
        assert browser["versions"][project] == version(project), (
            f"{project}: the website runs {browser['versions'][project]}, "
            f"the desktop has {version(project)}")

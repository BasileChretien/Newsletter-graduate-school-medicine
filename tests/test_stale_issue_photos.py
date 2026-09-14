"""Rebuilding an issue must not keep photos from earlier builds.

`assets/issue-N/` is the toolkit's own folder: `_build_pipeline` extracts
the Word file's photos into it and copies the drop-folder photos there.
Nothing ever emptied it, so a photo taken out of the Word file or out of
`drop-images/` stayed after the rebuild -- listed and counted by the
manifest, and pushed again by URL mode's publish step. A rebuild now
leaves the photos this build produced, and nothing else that is a photo.
"""

from __future__ import annotations

import base64
import re
import shutil
from pathlib import Path

import docx as docx_lib
import pytest

import build_newsletter as bn
from scripts.config import MERIDIAN_TEMPLATE
from scripts.manifest import load_manifest

# Smallest valid PNG (1x1). The drop folder checks magic bytes, so it has
# to be a real image.
_PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQ"
    "DwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)

# The template embeds three photos: word/media/image1.jpg .. image3.jpg.
_TEMPLATE_PHOTOS = {"image1.jpg", "image2.jpg", "image3.jpg"}


@pytest.fixture
def toolkit(tmp_path: Path, monkeypatch) -> Path:
    """A throwaway toolkit folder. dist/, assets/ and drop-images/ live
    under it, and photo URLs are resolved against it."""
    root = tmp_path / "toolkit"
    monkeypatch.setattr(bn, "PROJECT_ROOT", root)
    monkeypatch.setattr(bn, "DIST_DIR", root / "dist")
    monkeypatch.setattr(bn, "ASSETS_DIR", root / "assets")
    monkeypatch.setattr(bn, "DROP_DIR", root / "drop-images")
    return root


def _filled_template(tmp_path: Path) -> Path:
    """The shipped template with its masthead filled, so it validates."""
    d = docx_lib.Document(str(MERIDIAN_TEMPLATE))
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for old, new in (("VOL. XX", "VOL. 4"),
                                     ("ISSUE NO. XX", "ISSUE NO. 1"),
                                     ("MONTH YEAR", "JUNE 2026")):
                        if old in para.text:
                            for run in para.runs:
                                run.text = run.text.replace(old, new)
    src = tmp_path / "issue-3.docx"
    d.save(str(src))
    return src


def _build(toolkit: Path, docx_path: Path) -> Path:
    result = bn._build_pipeline(docx_path, issue=3, validate_remote=False)
    assert result.exit_code == 0
    return toolkit / "assets" / "issue-3"


def _photos(folder: Path) -> set[str]:
    return {f.name for f in folder.iterdir()
            if f.suffix.lower() in (".jpg", ".jpeg", ".png", ".gif")}


def test_a_drop_photo_taken_out_is_gone_after_the_rebuild(
    toolkit: Path, tmp_path: Path,
):
    drop = toolkit / "drop-images"
    drop.mkdir(parents=True)
    (drop / "s1_1_campus.png").write_bytes(_PNG_1X1)
    src = _filled_template(tmp_path)
    folder = _build(toolkit, src)
    assert "s1_1_campus.png" in _photos(folder)

    (drop / "s1_1_campus.png").unlink()
    folder = _build(toolkit, src)

    assert _photos(folder) == _TEMPLATE_PHOTOS
    manifest = load_manifest(folder)
    assert "s1_1_campus.png" not in manifest.files
    assert manifest.image_count == len(_TEMPLATE_PHOTOS)


def test_a_photo_the_word_file_no_longer_has_is_gone_after_the_rebuild(
    toolkit: Path, tmp_path: Path,
):
    """`image4.jpg` is what an earlier build of a longer Word file left
    behind. The rebuild removes it, and keeps every photo the email uses."""
    src = _filled_template(tmp_path)
    folder = _build(toolkit, src)
    shutil.copy(folder / "image3.jpg", folder / "image4.jpg")

    folder = _build(toolkit, src)

    assert _photos(folder) == _TEMPLATE_PHOTOS
    manifest = load_manifest(folder)
    assert "image4.jpg" not in manifest.files
    assert manifest.image_count == len(_TEMPLATE_PHOTOS)
    html = (toolkit / "dist" / "issue-3.html").read_text(encoding="utf-8")
    shown = set(re.findall(r"assets/issue-3/([\w.-]+)", html))
    assert shown, "the email should show the template's photos"
    assert all((folder / name).is_file() for name in shown)


def test_the_rebuild_keeps_what_is_not_a_photo(toolkit: Path, tmp_path: Path):
    src = _filled_template(tmp_path)
    folder = _build(toolkit, src)
    (folder / "notes.txt").write_text("kept", encoding="utf-8")

    folder = _build(toolkit, src)

    assert (folder / "notes.txt").read_text(encoding="utf-8") == "kept"
    assert (folder / "manifest.json").is_file()


def test_a_photo_already_there_under_another_case_is_not_deleted(
    toolkit: Path, tmp_path: Path,
):
    """On Windows and macOS `IMAGE1.JPG` and `image1.jpg` are one file:
    extracting `image1.jpg` writes into it and the old spelling stays on
    disk. Comparing names exactly would delete the photo just written,
    and the email would lose it."""
    folder = toolkit / "assets" / "issue-3"
    folder.mkdir(parents=True)
    (folder / "IMAGE1.JPG").write_bytes(b"old bytes")

    folder = _build(toolkit, _filled_template(tmp_path))

    assert (folder / "image1.jpg").is_file()


def test_a_word_file_that_cannot_be_read_leaves_the_photos_alone(
    toolkit: Path, tmp_path: Path,
):
    """Photos go only after this build has produced its own: a damaged
    Word file stops the build before that, and the last good build's
    photos stay."""
    folder = _build(toolkit, _filled_template(tmp_path))
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a Word file")

    result = bn._build_pipeline(broken, issue=3, validate_remote=False)

    assert result.exit_code == 1
    assert _photos(folder) == _TEMPLATE_PHOTOS
